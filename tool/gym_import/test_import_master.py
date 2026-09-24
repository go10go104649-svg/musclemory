import copy
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from import_master import read_master, prepare, normalized, sql_for

class ImportTests(unittest.TestCase):
    def setUp(self):
        self.tables = {'店舗':[dict(gym_id='s',gym_chain='Gym',gym_name="O'Brien 店",prefecture='県',municipality='市',machine_data_status='published')],
          '設備マスター':[dict(equipment_id='e',normalized_name='ダンベル',category='フリーウェイト',load_type='free_weight',needs_review=False)],
          '店舗設備':[dict(gym_id='s',equipment_id='e',raw_name='ダンベル',available=True,quantity=None)],'sheets':[]}
        self.catalog={'exercises':[dict(exerciseId='curl')]}
    def payload(self,t=None,m=None): return prepare(t or self.tables,'gym','Gym',m or [],self.catalog)
    def test_nulls_stable_keys_and_unmapped(self):
        a=self.payload(); b=self.payload()
        self.assertEqual(a,b)
        self.assertIsNone(a['gym_stores'][0]['station'])
        self.assertIsNone(a['gym_store_equipment'][0]['quantity'])
        self.assertIsNone(a['equipment'][0]['manufacturer'])
        self.assertEqual(a['equipment_exercise_mapping'],[])
        self.assertIn('on conflict (store_id,equipment_id)',sql_for(a))
        self.assertIn("O''Brien",sql_for(a))
    def test_duplicate_and_orphan_rejected(self):
        for table in ['店舗','設備マスター','店舗設備']:
            t=copy.deepcopy(self.tables); t[table].append(t[table][0].copy())
            with self.assertRaises(ValueError):self.payload(t)
        t=copy.deepcopy(self.tables);t['店舗設備'][0]['equipment_id']='missing'
        with self.assertRaises(ValueError):self.payload(t)
    def test_invalid_quantity_and_unknown_mapping_rejected(self):
        for value in [0,-1,1.2,'2',True]:
            t=copy.deepcopy(self.tables);t['店舗設備'][0]['quantity']=value
            with self.assertRaises(ValueError):self.payload(t)
        m=[dict(equipment_id='e',expected_name='ダンベル',expected_load_type='free_weight',exercise_ids=['missing'],rationale='test')]
        with self.assertRaises(ValueError):self.payload(m=m)
    def test_normalization_does_not_merge_ids(self):
        self.assertEqual(normalized(' ＡＢＣ　 24 '),'ABC 24')
        t=copy.deepcopy(self.tables);t['設備マスター'].append(dict(t['設備マスター'][0],equipment_id='e2'))
        self.assertEqual(len(self.payload(t)['equipment']),2)
    def test_tbar_equipment_variants_share_exercise_without_merging(self):
        mapping_path = Path(__file__).with_name('fitplace_mappings.json')
        all_mappings = json.loads(mapping_path.read_text())
        wanted = {
            'fp_eq_93cafce3bb2c',
            'fp_eq_aa7f2abd2795',
        }
        mappings = [m for m in all_mappings if m['equipment_id'] in wanted]
        self.assertEqual({m['equipment_id'] for m in mappings}, wanted)

        t = copy.deepcopy(self.tables)
        t['設備マスター'] = [
            dict(
                equipment_id='fp_eq_93cafce3bb2c',
                normalized_name='Tバーロー',
                category='フリーウェイト',
                load_type='not_specified',
                needs_review=False,
            ),
            dict(
                equipment_id='fp_eq_aa7f2abd2795',
                normalized_name='Tバーロー（プレートロード）',
                category='フリーウェイト',
                load_type='plate_loaded_explicit',
                needs_review=False,
            ),
        ]
        t['店舗設備'] = [
            dict(
                gym_id='s',
                equipment_id=e['equipment_id'],
                raw_name=e['normalized_name'],
                available=True,
                quantity=None,
            )
            for e in t['設備マスター']
        ]
        payload = prepare(
            t,
            'gym',
            'Gym',
            mappings,
            {'exercises': [dict(exerciseId='t_bar_row')]},
        )
        links = payload['equipment_exercise_mapping']
        self.assertEqual(len(links), 2)
        self.assertEqual({link['exercise_id'] for link in links}, {'t_bar_row'})
        self.assertEqual(len({link['equipment_id'] for link in links}), 2)

    def test_mapping_file_has_unique_equipment_and_valid_exercises(self):
        mapping_path = Path(__file__).with_name('fitplace_mappings.json')
        mappings = json.loads(mapping_path.read_text())
        catalog_path = Path(__file__).resolve().parents[1] / 'exercise_forms' / 'catalog.json'
        catalog = json.loads(catalog_path.read_text())
        selectable = {
            e['exerciseId']
            for e in catalog['exercises']
            if not e.get('canonicalExerciseId') and e.get('selectable', True)
        }
        equipment_ids = [m['equipment_id'] for m in mappings]
        self.assertEqual(len(equipment_ids), len(set(equipment_ids)))
        self.assertEqual(len(mappings), 185)
        self.assertEqual(sum(len(m['exercise_ids']) for m in mappings), 271)
        for mapping in mappings:
            self.assertTrue(mapping['exercise_ids'])
            self.assertEqual(len(mapping['exercise_ids']), len(set(mapping['exercise_ids'])))
            for exercise_id in mapping['exercise_ids']:
                self.assertIn(exercise_id, selectable)

    def test_read_xlsx_without_modification(self):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'input.xlsx';w=Workbook();w.remove(w.active)
            for name,rows in self.tables.items():
                if name=='sheets':continue
                s=w.create_sheet(name);s.append(list(rows[0]));s.append(list(rows[0].values()))
            w.save(path);before=path.read_bytes()
            t=read_master(path);self.assertEqual(len(self.payload(t)['gym_stores']),1)
            self.assertEqual(path.read_bytes(),before)
if __name__=='__main__': unittest.main()
