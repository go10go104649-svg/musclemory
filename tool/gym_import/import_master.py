"""Read a supplied workbook without modifying it; validate, then idempotently upsert.
Requires openpyxl. SQL execution uses the already-authenticated Supabase CLI.
No credentials or workbook contents are written to the repository.
"""
import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import unicodedata

ROOT = Path(__file__).resolve().parents[2]

def normalized(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value)).split())

def serial(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value

def read_master(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    tables = {}
    try:
        for name in ['店舗', '設備マスター', '店舗設備']:
            rows = iter(book[name].values)
            headers = next(rows)
            if len(headers) != len(set(headers)):
                raise ValueError(f'Duplicate columns: {name}')
            tables[name] = [{k: serial(v) for k, v in zip(headers, row)}
                            for row in rows if any(v is not None for v in row)]
        tables['sheets'] = book.sheetnames
    finally:
        book.close()
    return tables

def validate(tables):
    specs = [('店舗', ('gym_id',), ('gym_chain','gym_name','prefecture','municipality')),
             ('設備マスター', ('equipment_id',), ('normalized_name','category','load_type')),
             ('店舗設備', ('gym_id','equipment_id'), ('raw_name',))]
    for name, keys, required in specs:
        seen = set()
        for row in tables[name]:
            if any(not isinstance(row.get(k),str) or not row[k].strip() for k in keys+required):
                raise ValueError(f'Missing required field: {name}')
            key = tuple(row[k] for k in keys)
            if key in seen:
                raise ValueError(f'Duplicate key in {name}: {key}')
            seen.add(key)
    stores = {r['gym_id'] for r in tables['店舗']}
    equipment = {r['equipment_id'] for r in tables['設備マスター']}
    for row in tables['店舗設備']:
        if row['gym_id'] not in stores or row['equipment_id'] not in equipment:
            raise ValueError('Orphan store/equipment relation')
        q = row.get('quantity')
        if q is not None and (isinstance(q,bool) or not isinstance(q,(int,float)) or q <= 0 or int(q)!=q):
            raise ValueError('Invalid explicit quantity')
        if not isinstance(row.get('available'),bool):
            raise ValueError('Invalid availability')

def prepare(tables, chain_id, chain_name, mappings, catalog):
    validate(tables)
    if not chain_id or any(r['gym_chain'] != chain_name for r in tables['店舗']):
        raise ValueError('Chain must match the workbook')
    prefix = lambda key: f'{chain_id}:{key}'
    stores = [dict(id=prefix(r['gym_id']),chain_id=chain_id,source_id=r['gym_id'],name=r['gym_name'],
      prefecture=r.get('prefecture'),city=r.get('municipality'),address=r.get('address_raw'),
      station=r.get('station'),official_url=r.get('official_url'),equipment_status=r['machine_data_status'],source=r)
      for r in tables['店舗']]
    equipment = [dict(id=prefix(r['equipment_id']),name=r['normalized_name'],normalized_name=normalized(r['normalized_name']),
      category=r['category'],load_type=r['load_type'],manufacturer=r.get('manufacturer_id'),model=r.get('model_name'),
      needs_review=bool(r.get('needs_review')),source=r) for r in tables['設備マスター']]
    relations = [dict(store_id=prefix(r['gym_id']),equipment_id=prefix(r['equipment_id']),quantity=int(r['quantity']) if r.get('quantity') is not None else None,
      available=r['available'],raw_name=r['raw_name'],source_url=r.get('source_url'),checked_at=r.get('checked_at'),source=r)
      for r in tables['店舗設備']]
    by_id = {e['exerciseId']:e for e in catalog['exercises']}
    source_equipment = {r['equipment_id']:r for r in tables['設備マスター']}
    links = []
    for m in mappings:
        row = source_equipment.get(m['equipment_id'])
        if row is None or row['normalized_name'] != m['expected_name'] or row['load_type'] != m['expected_load_type']:
            raise ValueError(f'Mapping source changed: {m["equipment_id"]}')
        if row['needs_review']:
            raise ValueError('Unreviewed equipment cannot be automatically mapped')
        for id in m['exercise_ids']:
            if id not in by_id or by_id[id].get('canonicalExerciseId') or not by_id[id].get('selectable',True):
                raise ValueError(f'Invalid exercise ID: {id}')
            links.append(dict(equipment_id=prefix(row['equipment_id']),exercise_id=id,rationale=m['rationale']))
    if len({(r['equipment_id'],r['exercise_id']) for r in links}) != len(links):
        raise ValueError('Duplicate exercise mapping')
    return {'gym_chains':[dict(id=chain_id,name=chain_name)],'gym_stores':stores,'equipment':equipment,
            'gym_store_equipment':relations,'equipment_exercise_mapping':links}

TYPES = {'source':'jsonb','needs_review':'boolean','available':'boolean','quantity':'integer','checked_at':'timestamptz'}
KEYS = {'gym_chains':['id'],'gym_stores':['id'],'equipment':['id'],
        'gym_store_equipment':['store_id','equipment_id'],'equipment_exercise_mapping':['equipment_id','exercise_id']}

def sql_for(payload):
    statements = ['begin;']
    for table, rows in payload.items():
        if not rows: continue
        columns = list(rows[0]); keys=KEYS[table]
        updates = ','.join(f'{c}=excluded.{c}' for c in columns if c not in keys)
        data = json.dumps(rows,ensure_ascii=False).replace("'", "''")
        definitions = ','.join(f'{c} {TYPES.get(c,"text")}' for c in columns)
        statements.append(f"insert into public.{table} ({','.join(columns)}) select {','.join(columns)} from jsonb_to_recordset('{data}'::jsonb) as r({definitions}) on conflict ({','.join(keys)}) do update set {updates};")
    statements.append('commit;')
    return '\n'.join(statements)

def summary(tables,payload):
    mapped={r['equipment_id'] for r in payload['equipment_exercise_mapping']}
    return dict(sheets=tables['sheets'],**{k:len(v) for k,v in payload.items()},mapped_equipment=len(mapped),
      unmapped_equipment=len(payload['equipment'])-len(mapped),
      stores_with_equipment=len({r['store_id'] for r in payload['gym_store_equipment']}),
      equipment_status=dict(Counter(r['machine_data_status'] for r in tables['店舗'])),
      null_station=sum(r['station'] is None for r in payload['gym_stores']),
      null_manufacturer=sum(r['manufacturer'] is None for r in payload['equipment']),
      null_model=sum(r['model'] is None for r in payload['equipment']),
      null_quantity=sum(r['quantity'] is None for r in payload['gym_store_equipment']))

def main():
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True)
    p.add_argument('--chain-id',required=True);p.add_argument('--chain-name',required=True)
    p.add_argument('--mapping',type=Path,required=True);p.add_argument('--apply',action='store_true')
    args=p.parse_args()
    before=hashlib.sha256(args.input.read_bytes()).hexdigest()
    tables=read_master(args.input)
    payload=prepare(tables,args.chain_id,args.chain_name,json.loads(args.mapping.read_text()),
      json.loads((ROOT/'tool/exercise_forms/catalog.json').read_text()))
    report=summary(tables,payload);report['input_sha256']=before
    if args.apply:
        # Temporary SQL is not retained or added to Git; argv never contains credentials.
        with tempfile.TemporaryDirectory(prefix='musclemory-gym-') as temp:
            path=Path(temp)/'import.sql';path.write_text(sql_for(payload))
            result=subprocess.run(['supabase','db','query','--linked','--file',str(path),'--output','json'],cwd=ROOT,capture_output=True,text=True)
            if result.returncode:
                raise RuntimeError('Supabase import failed: '+result.stderr[:500])
        report['applied']=True
    assert hashlib.sha256(args.input.read_bytes()).hexdigest()==before
    print(json.dumps(report,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
